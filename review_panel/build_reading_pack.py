#!/usr/bin/env python3
"""Build an easy-to-read Word review pack from completed CALON-C panel Markdown outputs."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RUN_ROOT = ROOT / "outputs/manuscript_2026-08-16/multi_agent_review_2026-08-17"
GOOGLE = RUN_ROOT / "google-gemini-3.6-flash-high"
CURSOR = RUN_ROOT / "cursor-three-model-four-seat"
PACK = Path.home() / "Documents/CALON/CALON_REVIEW_READING_PACK_2026-08-17"
SOURCE_MD = PACK / "source_markdown"

REPORTS = [
    ("01_GOOGLE_FINAL_TEACHING_SYNTHESIS", "Google Gemini — Final teaching synthesis", GOOGLE / "final/TEACHING_SYNTHESIS.md", "COMPLETE"),
    ("02_GOOGLE_BIOSTATISTICIAN", "Google Gemini — Biostatistician, blind review", GOOGLE / "round-1-independent/biostatistician.md", "COMPLETE"),
    ("03_GOOGLE_CARDIOLOGIST", "Google Gemini — Cardiologist, blind review", GOOGLE / "round-1-independent/cardiologist.md", "COMPLETE"),
    ("04_GOOGLE_LIPID_MEDICINE", "Google Gemini — Lipid-medicine specialist, blind review", GOOGLE / "round-1-independent/lipid_medicine.md", "COMPLETE"),
    ("05_GOOGLE_SENIOR_EDITOR", "Google Gemini — Senior Editor-in-Chief, blind review", GOOGLE / "round-1-independent/senior_editor_in_chief.md", "COMPLETE"),
    ("06_GOOGLE_DEBATE_BIOSTATISTICIAN", "Google Gemini — Biostatistician, debate round", GOOGLE / "round-2-debate/biostatistician.md", "COMPLETE"),
    ("07_GOOGLE_DEBATE_CARDIOLOGIST", "Google Gemini — Cardiologist, debate round", GOOGLE / "round-2-debate/cardiologist.md", "COMPLETE"),
    ("08_GOOGLE_DEBATE_LIPID_MEDICINE", "Google Gemini — Lipid-medicine specialist, debate round", GOOGLE / "round-2-debate/lipid_medicine.md", "COMPLETE"),
    ("09_GOOGLE_DEBATE_SENIOR_EDITOR", "Google Gemini — Senior Editor-in-Chief, debate round", GOOGLE / "round-2-debate/senior_editor_in_chief.md", "COMPLETE"),
    ("10_CURSOR_KIMI_ROUND1", "Cursor Kimi — Four-seat internal panel, blind round", CURSOR / "round-1-independent/kimi.md", "COMPLETE ROUND 1; CURSOR PIPELINE CONTINUES"),
    ("11_CURSOR_GROK_ROUND1", "Cursor Grok — Four-seat internal panel, blind round", CURSOR / "round-1-independent/grok.md", "COMPLETE ROUND 1; CURSOR PIPELINE CONTINUES"),
    ("12_CURSOR_CLAUDE_ROUND1", "Cursor Claude — Four-seat internal panel, blind round", CURSOR / "round-1-independent/claude.md", "COMPLETE ROUND 1; CURSOR PIPELINE CONTINUES"),
    ("13_CURSOR_GROK_DEBATE", "Cursor Grok — Cross-model debate response", CURSOR / "round-2-cross-model-debate/grok.md", "COMPLETE; KIMI/CLAUDE DEBATE AND FINAL SYNTHESIS NOT COMPLETED"),
]


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"[Image: \1]", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("~~", "")
    return text.strip()


def configure(doc: Document, title: str) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    for level in range(1, 5):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor(31, 78, 121)
    for section in doc.sections:
        section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75); section.right_margin = Inches(0.75)
        header = section.header.paragraphs[0]
        header.text = "CALON-C manuscript review · generated review pack"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(8); header.runs[0].font.color.rgb = RGBColor(90, 90, 90)
        footer = section.footer.paragraphs[0]
        footer.text = "CALON-C manuscript review"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(8); footer.runs[0].font.color.rgb = RGBColor(110, 110, 110)
    doc.core_properties.title = title
    doc.core_properties.author = "CALON multi-agent review panel"


def add_title_page(doc: Document, title: str, subtitle: str, status: str) -> None:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title).bold = True; p.runs[0].font.size = Pt(24); p.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(subtitle).italic = True; p.runs[0].font.size = Pt(13)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(status); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(156, 87, 0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Prepared 17 August 2026").font.size = Pt(10)
    doc.add_paragraph()


def parse_table(lines: list[str], start: int):
    if start + 1 >= len(lines) or "|" not in lines[start]: return None
    separator = lines[start + 1].strip()
    if not re.match(r"^\s*\|?\s*:?-{3,}", separator) or "|" not in separator: return None
    rows = []
    i = start
    while i < len(lines) and "|" in lines[i] and lines[i].strip():
        rows.append([clean_inline(c) for c in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) < 2: return None
    rows.pop(1)
    return rows, i


def add_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines(); i = 0; in_code = False; code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True; code_lines = []
            else:
                p = doc.add_paragraph(style="No Spacing")
                r = p.add_run("\n".join(code_lines)); r.font.name = "Courier New"; r.font.size = Pt(8)
                p.paragraph_format.space_after = Pt(6); in_code = False
            i += 1; continue
        if in_code:
            code_lines.append(line); i += 1; continue
        table_data = parse_table(lines, i)
        if table_data:
            rows, next_i = table_data
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Light Shading Accent 1"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(width):
                    cell = table.cell(ri, ci); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    cell.text = row[ci] if ci < len(row) else ""
                    for run in cell.paragraphs[0].runs: run.font.size = Pt(8.5); run.bold = ri == 0
            i = next_i; continue
        stripped = line.strip()
        if not stripped or stripped in ("---", "***", "___"):
            if stripped == "": doc.add_paragraph()
            i += 1; continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            doc.add_heading(clean_inline(heading.group(2)), level=min(len(heading.group(1)), 4)); i += 1; continue
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        if bullet:
            doc.add_paragraph(clean_inline(bullet.group(1)), style="List Bullet"); i += 1; continue
        number = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if number:
            doc.add_paragraph(clean_inline(number.group(1)), style="List Number"); i += 1; continue
        if stripped.startswith(">"):
            p = doc.add_paragraph(clean_inline(stripped.lstrip("> ")))
            p.paragraph_format.left_indent = Inches(0.3); p.runs[0].italic = True
            i += 1; continue
        p = doc.add_paragraph(clean_inline(stripped))
        p.paragraph_format.space_after = Pt(4)
        i += 1


def save_doc(path: Path, title: str, subtitle: str, status: str, sections: list[tuple[str, str]]) -> None:
    doc = Document(); configure(doc, title); add_title_page(doc, title, subtitle, status)
    for index, (section_title, markdown) in enumerate(sections):
        heading = doc.add_heading(section_title, level=1)
        if index:
            heading.paragraph_format.page_break_before = True
        add_markdown(doc, markdown)
    doc.save(path)


def main() -> None:
    PACK.mkdir(parents=True, exist_ok=True); SOURCE_MD.mkdir(exist_ok=True)
    available = [(stem, title, path, status) for stem, title, path, status in REPORTS if path.exists()]
    if not available: raise RuntimeError("No completed review Markdown files found")
    for stem, title, path, status in available:
        text = path.read_text(encoding="utf-8")
        save_doc(PACK / f"{stem}.docx", title, "CALON-C paragraph-by-paragraph manuscript review", status, [(title, text)])
        shutil.copy2(path, SOURCE_MD / f"{stem}.md")

    # One combined document for direct searching and comparison in Word.
    combined_sections = [(f"{stem} · {title} · {status}", path.read_text(encoding="utf-8")) for stem, title, path, status in available]
    save_doc(
        PACK / "00_MASTER_COMPARISON_ALL_FINISHED_REVIEWS.docx",
        "CALON-C: all finished panel reviews",
        "Google specialist reviews and debates, plus completed Cursor blind reviews",
        "READ AS PANEL OPINION — REQUIRES SCIENTIFIC ADJUDICATION",
        combined_sections,
    )

    # Start-here status and file index.
    doc = Document(); configure(doc, "CALON-C review reading pack — start here")
    add_title_page(doc, "CALON-C review reading pack", "Finished reports arranged for rapid reading and comparison", "GOOGLE COMPLETE · CURSOR PARTIALLY COMPLETE")
    doc.add_heading("How to use this folder", level=1)
    for item in [
        "Start with 01_GOOGLE_FINAL_TEACHING_SYNTHESIS.docx for the condensed editorial view.",
        "Use 00_MASTER_COMPARISON_ALL_FINISHED_REVIEWS.docx to search and compare every report in one Word file.",
        "Read 02–05 for the four independent Google specialist reviews.",
        "Read 06–09 for the four Google debate-round responses.",
        "Read 10–12 for the completed Cursor Kimi, Grok and Claude blind-round reports, and 13 for Grok's completed cross-model debate response. Kimi's and Claude's debate responses and the final Cursor teaching synthesis were not completed.",
        "Treat numerical rejection probabilities and new analysis recommendations as panel opinion until checked against the live code, data and methods literature.",
    ]: doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("File index", level=1)
    table = doc.add_table(rows=1, cols=3); table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ["Word file", "Report", "Status"]): cell.text = value
    for stem, title, path, status in available:
        row = table.add_row().cells; row[0].text = f"{stem}.docx"; row[1].text = title; row[2].text = status
    doc.add_heading("Current source locations", level=1)
    doc.add_paragraph(str(GOOGLE)); doc.add_paragraph(str(CURSOR))
    doc.save(PACK / "00_START_HERE.docx")

    readme = [
        "# CALON-C review reading pack", "", "Built: 17 August 2026", "",
        "Open `00_START_HERE.docx` first, then `00_MASTER_COMPARISON_ALL_FINISHED_REVIEWS.docx`.", "",
        "## Status", "", "- Google Gemini four-specialist panel: complete.",
        "- Cursor panel: Kimi, Grok and Claude blind reports and Grok's debate response are included; the remaining Cursor debate responses and final synthesis were not completed.", "",
        "## Caution", "", "These are AI-panel opinions, not an adjudicated evidence set. Verify every proposed new analysis and citation before changing the manuscript.", "",
    ]
    (PACK / "README.md").write_text("\n".join(readme), encoding="utf-8")
    print(f"PACK={PACK}")
    print(f"DOCX_COUNT={len(list(PACK.glob('*.docx')))}")
    for path in sorted(PACK.glob("*.docx")): print(path.name)


if __name__ == "__main__":
    main()
