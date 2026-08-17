#!/usr/bin/env python3
"""Build polished CALON-N manuscript and supplement DOCX files."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("/Users/nader85/Documents/CALON/calon_discordance_model_2026_08_09")
BLUE = "2E5E7E"
DARK = "19384D"
MUTED = "667784"
LIGHT = "EAF0F4"
GRID = "B8C5CE"


def set_font(run, name="Calibri", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)


def page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure(doc: Document, landscape=False, header_label="Original Research"):
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)
        sec.left_margin = sec.right_margin = Inches(0.7)
        sec.top_margin = sec.bottom_margin = Inches(0.75)
    else:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(1)
        sec.top_margin = sec.bottom_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ]:
        st = styles[style_name]; st.font.name = "Calibri"; st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color); st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Number"]:
        st = styles[name]; st.font.name = "Calibri"; st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.194)
        st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.208

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run(f"CALON-N  |  {header_label}"), size=8.5, color=MUTED)
    page_field(sec.footer.paragraphs[0])


def cover(doc: Document, title: str, subtitle: str, document_label: str):
    for _ in range(5):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    set_font(p.add_run(document_label.upper()), size=10, bold=True, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    set_font(p.add_run(title), size=23, bold=True, color=DARK)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(subtitle), size=13, italic=True, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("Target journal: Journal of Clinical Lipidology"), size=10.5, color=MUTED)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Analysis freeze: 9 August 2026"), size=10, color=MUTED)
    doc.add_page_break()


def add_inline(paragraph, text: str, size=11):
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=size, bold=True)
        elif token.startswith("`"):
            set_font(paragraph.add_run(token[1:-1]), name="Courier New", size=max(size-1, 8.5), color=DARK)
        else:
            set_font(paragraph.add_run(token[1:-1]), size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]), size=size)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar"); tc_pr.append(mar)
    for edge, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = mar.find(qn(f"w:{edge}"))
        if node is None: node = OxmlElement(f"w:{edge}"); mar.append(node)
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")


def table_geometry(table, widths, indent=120):
    total = sum(widths); table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag, attrs in [
        ("w:tblW", {"w:w": str(total), "w:type": "dxa"}),
        ("w:tblInd", {"w:w": str(indent), "w:type": "dxa"}),
        ("w:tblLayout", {"w:type": "fixed"}),
    ]:
        old = tbl_pr.find(qn(tag))
        if old is not None: tbl_pr.remove(old)
        node = OxmlElement(tag)
        for k, v in attrs.items(): node.set(qn(k), v)
        tbl_pr.append(node)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr(); tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None: tc_w = OxmlElement("w:tcW"); tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width)); tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def widths_for(rows, total):
    n = len(rows[0]); lengths = []
    for j in range(n):
        lengths.append(max(8, min(42, max(len(str(r[j])) for r in rows))))
    raw = np_array = [x / sum(lengths) for x in lengths]
    minimum = 650 if n <= 6 else 550
    widths = [max(minimum, int(total * x)) for x in raw]
    scale = total / sum(widths); widths = [int(x * scale) for x in widths]
    widths[-1] += total - sum(widths)
    return widths


def add_table(doc, rows, landscape=False):
    cols = len(rows[0]); table = doc.add_table(rows=len(rows), cols=cols)
    total = 13824 if landscape else 9360
    widths = widths_for(rows, total)
    table_geometry(table, widths, indent=120)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j); cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if j > 0 and re.fullmatch(r"[+−\-]?[0-9.,%±–—() /]+", value.strip()):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value, size=8.4 if cols >= 6 else 9)
            if i == 0:
                shade(cell, LIGHT)
                for run in p.runs: run.bold = True; run.font.color.rgb = RGBColor.from_string(DARK)
    tr_pr = table.rows[0]._tr.get_or_add_trPr(); header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true"); tr_pr.append(header)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells): continue
        rows.append(cells)
    return rows


def add_markdown(doc, path: Path, landscape=False, add_figures=False):
    lines = path.read_text(encoding="utf-8").splitlines()
    paragraph = []
    code = False; code_lines = []

    def flush():
        nonlocal paragraph
        if paragraph:
            text = " ".join(x.strip() for x in paragraph)
            p = doc.add_paragraph(); add_inline(p, text)
            paragraph = []

    i = 1  # title handled by cover
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush()
            if code:
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25)
                p.paragraph_format.right_indent = Inches(.25); p.paragraph_format.space_after = Pt(10)
                shade_para = OxmlElement("w:shd"); shade_para.set(qn("w:fill"), "F3F5F7")
                p._p.get_or_add_pPr().append(shade_para)
                set_font(p.add_run("\n".join(code_lines)), name="Courier New", size=9, color=DARK)
                code_lines = []; code = False
            else: code = True
            i += 1; continue
        if code:
            code_lines.append(line); i += 1; continue
        if line.startswith("|"):
            flush(); block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, parse_table(block), landscape=landscape); continue
        if not line.strip():
            flush(); i += 1; continue
        if line.startswith("### "):
            flush(); doc.add_heading(line[4:].strip(), level=2); i += 1; continue
        if line.startswith("## "):
            flush(); doc.add_heading(line[3:].strip(), level=1); i += 1; continue
        if line.startswith("# "):
            flush(); doc.add_heading(line[2:].strip(), level=1); i += 1; continue
        if line.startswith("- "):
            flush(); item = [line[2:].strip()]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3} |\||- |\d+\. |```)", lines[i]):
                item.append(lines[i].strip()); i += 1
            p = doc.add_paragraph(style="List Bullet"); add_inline(p, " " + " ".join(item)); continue
        if re.match(r"^\d+\. ", line):
            flush(); number = re.match(r"^(\d+)\. ", line).group(1)
            item = [re.sub(r"^\d+\. ", "", line).strip()]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3} |\||- |\d+\. |```)", lines[i]):
                item.append(lines[i].strip()); i += 1
            # Preserve the explicit source number. Word's List Number style can
            # continue numbering across unrelated lists (for example, the five
            # transport mechanisms followed by the reference list).
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.22)
            p.paragraph_format.first_line_indent = Inches(-.22)
            add_inline(p, f"{number}. " + " ".join(item)); continue
        paragraph.append(line); i += 1
    flush()
    if add_figures:
        doc.add_page_break(); doc.add_heading("Figures", level=1)
        for image, caption in [
            (ROOT / "figures" / "scratch_transport_auc.png", "Figure 1. Reciprocal transport AUC for locked raw-variable candidates."),
            (ROOT / "figures" / "scratch_external_calibration.png", "Figure 2. CALON-N calibration in reciprocal transport without target recalibration."),
        ]:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            shape = p.add_run().add_picture(str(image), width=Inches(6.2))
            shape._inline.docPr.set("title", caption.split(".", 1)[0])
            shape._inline.docPr.set("descr", caption)
            c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(c.add_run(caption), size=9.5, italic=True, color=MUTED)


def build(source, output, title, subtitle, landscape=False, figures=False, header_label="Original Research"):
    doc = Document(); configure(doc, landscape=landscape, header_label=header_label)
    cover(doc, title, subtitle, header_label)
    add_markdown(doc, source, landscape=landscape, add_figures=figures)
    props = doc.core_properties
    props.title = title; props.subject = "CALON-N manuscript package"; props.author = "CALON investigators"
    props.keywords = "familial hypercholesterolaemia; apoB; LDL; ASCVD; model"
    doc.save(output)


def main():
    build(
        ROOT / "manuscript" / "MANUSCRIPT.md", ROOT / "manuscript" / "CALON_N_MANUSCRIPT.docx",
        "Apolipoprotein B–LDL discordance identifies established cardiovascular disease in genetically defined familial hypercholesterolaemia",
        "Absolute risk does not transport between clinic and biobank", figures=True,
    )
    build(
        ROOT / "manuscript" / "SUPPLEMENT.md", ROOT / "manuscript" / "CALON_N_SUPPLEMENT.docx",
        "CALON-N online supplement", "Methods, cohort audit, candidate performance, and sensitivity analyses",
        landscape=True, figures=False, header_label="Online Supplement",
    )
    build(
        ROOT / "manuscript" / "ANALYTICAL_METHODS_RESULTS_COMPENDIUM.md",
        ROOT / "manuscript" / "CALON_N_ANALYTICAL_METHODS_RESULTS_COMPENDIUM.docx",
        "CALON-N full analytical methods and results compendium",
        "Complete cohort construction, variables, modelling, performance, sensitivities, and QC",
        landscape=True, figures=False, header_label="Analytical Compendium",
    )
    build(
        ROOT / "qc" / "TRIPOD_STROBE_PROBAST_AUDIT.md",
        ROOT / "manuscript" / "CALON_N_TRIPOD_STROBE_PROBAST_AUDIT.docx",
        "CALON-N reporting and risk-of-bias audit",
        "Full TRIPOD+AI and STROBE mapping with PROBAST judgement",
        landscape=True, figures=False, header_label="Reporting Audit",
    )


if __name__ == "__main__":
    main()
